"""
General-purpose Streamlit app for asking OpenAI and Claude the same question,
optionally with uploaded attachments, then letting the two model outputs review
one another before producing a final synthesis.

Local run:
    pip install -r requirements.txt
    streamlit run multi_model_chat_app.py

GitHub + Streamlit Community Cloud deployment:
    1. Put this file in a GitHub repository.
    2. Add requirements.txt to the same repository.
    3. Deploy the repository from Streamlit Community Cloud.
    4. Add API keys in Streamlit Cloud Secrets, not in GitHub.

Streamlit Secrets format:
    OPENAI_API_KEY = "your_openai_key"
    ANTHROPIC_API_KEY = "your_anthropic_key"

For local development, you can either use Streamlit secrets at
.streamlit/secrets.toml or environment variables.
"""

from __future__ import annotations

import os
from dataclasses import dataclass
from io import BytesIO
from typing import Optional

import pandas as pd
import streamlit as st
from anthropic import Anthropic
from openai import OpenAI
from pypdf import PdfReader
from docx import Document


# -----------------------------
# Configuration
# -----------------------------

# These are safe, relatively economical defaults.
# You can change them in the Streamlit sidebar.
DEFAULT_OPENAI_MODEL = "gpt-4.1-mini"
DEFAULT_CLAUDE_MODEL = "claude-3-5-haiku-latest"

# Cost controls.
# The app extracts text from uploads and sends that text to both APIs.
# Lower these numbers if you want to reduce API cost.
MAX_ATTACHMENT_CHARS = 60_000
MAX_OUTPUT_TOKENS = 1800


@dataclass
class ModelResult:
    model_name: str
    answer: str


# -----------------------------
# Secret handling
# -----------------------------

def get_secret(name: str) -> Optional[str]:
    """Read secrets from Streamlit Secrets first, then environment variables.

    Streamlit Community Cloud uses st.secrets.
    Local development can use either .streamlit/secrets.toml or environment variables.
    """
    try:
        if name in st.secrets:
            return st.secrets[name]
    except Exception:
        pass
    return os.getenv(name)


# -----------------------------
# File extraction helpers
# -----------------------------

def extract_text_from_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(f"\n--- PDF page {i} ---\n{text}")
    return "\n".join(pages).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = Document(BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    table_text = []
    for table_index, table in enumerate(doc.tables, start=1):
        table_text.append(f"\n--- DOCX table {table_index} ---")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            table_text.append(" | ".join(cells))

    return "\n".join(paragraphs + table_text).strip()


def extract_text_from_spreadsheet(file_bytes: bytes, filename: str) -> str:
    if filename.lower().endswith(".csv"):
        df = pd.read_csv(BytesIO(file_bytes))
    else:
        df = pd.read_excel(BytesIO(file_bytes))
    return df.to_csv(index=False)


def extract_text_from_plain_text(file_bytes: bytes) -> str:
    for encoding in ["utf-8", "utf-16", "latin-1"]:
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore")


def extract_uploaded_file_text(uploaded_file) -> str:
    filename = uploaded_file.name
    suffix = filename.lower().split(".")[-1]
    file_bytes = uploaded_file.getvalue()

    try:
        if suffix == "pdf":
            text = extract_text_from_pdf(file_bytes)
        elif suffix == "docx":
            text = extract_text_from_docx(file_bytes)
        elif suffix in ["csv", "xlsx", "xls"]:
            text = extract_text_from_spreadsheet(file_bytes, filename)
        elif suffix in ["txt", "md", "json", "xml", "html", "py", "r", "sql", "js", "ts", "css"]:
            text = extract_text_from_plain_text(file_bytes)
        else:
            text = f"[Unsupported file type for {filename}. Please convert it to PDF, DOCX, TXT, CSV, or XLSX.]"
    except Exception as exc:
        text = f"[Could not extract text from {filename}. Error: {exc}]"

    if len(text) > MAX_ATTACHMENT_CHARS:
        text = text[:MAX_ATTACHMENT_CHARS] + "\n\n[Attachment text was truncated because it is very long.]"

    return f"\n\n===== ATTACHMENT: {filename} =====\n{text}\n===== END ATTACHMENT: {filename} =====\n"


def build_attachment_context(uploaded_files) -> str:
    if not uploaded_files:
        return ""
    return "\n".join(extract_uploaded_file_text(file) for file in uploaded_files)


# -----------------------------
# Prompt construction
# -----------------------------

def build_initial_prompt(user_question: str, user_context: str, attachment_context: str) -> str:
    return f"""
You are one of two AI assistants helping the user with a general-purpose task.

Task from user:
{user_question}

User's additional context, preferences, rubric, or opinion:
{user_context if user_context.strip() else "[No additional user context provided.]"}

Attachments extracted as text:
{attachment_context if attachment_context.strip() else "[No attachments provided.]"}

Instructions:
1. Answer the user's task directly.
2. Use the attachments when relevant.
3. Be explicit about assumptions, uncertainty, and limitations.
4. Structure your answer clearly.
5. Do not claim to have seen images, layout, or visual details unless they are represented in the extracted text.
""".strip()


def build_cross_review_prompt(
    user_question: str,
    user_context: str,
    attachment_context: str,
    own_initial_answer: str,
    other_model_name: str,
    other_initial_answer: str,
) -> str:
    return f"""
The user asked:
{user_question}

User's additional context, preferences, rubric, or opinion:
{user_context if user_context.strip() else "[No additional user context provided.]"}

Attachments extracted as text:
{attachment_context if attachment_context.strip() else "[No attachments provided.]"}

Your initial answer was:
{own_initial_answer}

The other model, {other_model_name}, answered:
{other_initial_answer}

Now revise your answer after considering the other model's reasoning.

Instructions:
1. Identify where the other model made a useful point.
2. Identify where you disagree or where the other model may be incomplete.
3. Produce a revised answer for the user.
4. Keep the final revised answer practical and clear.
""".strip()


def build_final_synthesis_prompt(
    user_question: str,
    user_context: str,
    openai_initial: str,
    claude_initial: str,
    openai_revised: str,
    claude_revised: str,
) -> str:
    return f"""
The user asked:
{user_question}

User's additional context, preferences, rubric, or opinion:
{user_context if user_context.strip() else "[No additional user context provided.]"}

OpenAI initial answer:
{openai_initial}

Claude initial answer:
{claude_initial}

OpenAI revised answer after seeing Claude:
{openai_revised}

Claude revised answer after seeing OpenAI:
{claude_revised}

Create one final synthesized answer for the user.

Instructions:
1. Combine the strongest points from both models.
2. Resolve disagreements clearly.
3. Mention uncertainty where needed.
4. Do not invent facts not supported by the prompt, user context, or extracted attachments.
5. Make the final output directly usable.
""".strip()


# -----------------------------
# Model call wrappers
# -----------------------------

def call_openai(prompt: str, model: str) -> str:
    api_key = get_secret("OPENAI_API_KEY")
    if not api_key:
        return "[Missing OPENAI_API_KEY. Add it to Streamlit Secrets or your local environment variables.]"

    try:
        client = OpenAI(api_key=api_key)
        response = client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": "You are a careful, practical assistant."},
                {"role": "user", "content": prompt},
            ],
            temperature=0.2,
            max_tokens=MAX_OUTPUT_TOKENS,
        )
        return response.choices[0].message.content or ""
    except Exception as exc:
        return f"[OpenAI API error: {exc}]"


def call_claude(prompt: str, model: str) -> str:
    api_key = get_secret("ANTHROPIC_API_KEY")
    if not api_key:
        return "[Missing ANTHROPIC_API_KEY. Add it to Streamlit Secrets or your local environment variables.]"

    try:
        client = Anthropic(api_key=api_key)
        message = client.messages.create(
            model=model,
            max_tokens=MAX_OUTPUT_TOKENS,
            temperature=0.2,
            system="You are a careful, practical assistant.",
            messages=[{"role": "user", "content": prompt}],
        )

        parts = []
        for block in message.content:
            if getattr(block, "type", None) == "text":
                parts.append(block.text)
        return "\n".join(parts).strip()
    except Exception as exc:
        return f"[Claude API error: {exc}]"


# -----------------------------
# Streamlit UI
# -----------------------------

st.set_page_config(page_title="Multi-Model Chat: OpenAI + Claude", layout="wide")

st.title("Multi-Model Chat: OpenAI + Claude")
st.caption("Ask one question, upload files once, compare both models, cross-review, and synthesize a final answer.")

with st.sidebar:
    st.header("Settings")
    openai_model = st.text_input("OpenAI model", value=DEFAULT_OPENAI_MODEL)
    claude_model = st.text_input("Claude model", value=DEFAULT_CLAUDE_MODEL)

    st.markdown("---")
    st.subheader("Workflow")
    enable_openai = st.checkbox("Use OpenAI", value=True)
    enable_claude = st.checkbox("Use Claude", value=True)
    enable_cross_review = st.checkbox("Let both models review each other", value=True)
    enable_final_synthesis = st.checkbox("Create final synthesized answer", value=True)

    st.markdown("---")
    st.subheader("API key status")
    st.write("OpenAI key:", "Found" if get_secret("OPENAI_API_KEY") else "Missing")
    st.write("Claude key:", "Found" if get_secret("ANTHROPIC_API_KEY") else "Missing")

    st.markdown("---")
    st.subheader("Cost control")
    st.write(f"Maximum attachment characters per file: {MAX_ATTACHMENT_CHARS:,}")
    st.write(f"Maximum output tokens per model call: {MAX_OUTPUT_TOKENS:,}")

user_question = st.text_area(
    "Your question or task",
    height=160,
    placeholder="Example: Compare these two policy drafts and recommend which one is stronger."
)

user_context = st.text_area(
    "Your own opinion, preference, rubric, or additional instructions — optional",
    height=120,
    placeholder="Example: I care most about clarity, fairness, and evidence. Please be conservative when making claims."
)

uploaded_files = st.file_uploader(
    "Upload attachments once",
    type=["pdf", "docx", "txt", "md", "csv", "xlsx", "xls", "json", "xml", "html", "py", "r", "sql", "js", "ts", "css"],
    accept_multiple_files=True,
)

run_button = st.button("Ask selected models", type="primary")

if run_button:
    if not user_question.strip():
        st.error("Please enter a question or task.")
        st.stop()

    if not enable_openai and not enable_claude:
        st.error("Please select at least one model in the sidebar.")
        st.stop()

    with st.spinner("Extracting attachment text..."):
        attachment_context = build_attachment_context(uploaded_files)

    initial_prompt = build_initial_prompt(user_question, user_context, attachment_context)

    st.subheader("1. Initial answers")

    openai_initial = ""
    claude_initial = ""

    if enable_openai and enable_claude:
        col1, col2 = st.columns(2)

        with st.spinner("Calling OpenAI and Claude..."):
            openai_initial = call_openai(initial_prompt, openai_model)
            claude_initial = call_claude(initial_prompt, claude_model)

        with col1:
            st.markdown("### OpenAI initial answer")
            st.markdown(openai_initial)

        with col2:
            st.markdown("### Claude initial answer")
            st.markdown(claude_initial)

    elif enable_openai:
        with st.spinner("Calling OpenAI..."):
            openai_initial = call_openai(initial_prompt, openai_model)
        st.markdown("### OpenAI initial answer")
        st.markdown(openai_initial)

    elif enable_claude:
        with st.spinner("Calling Claude..."):
            claude_initial = call_claude(initial_prompt, claude_model)
        st.markdown("### Claude initial answer")
        st.markdown(claude_initial)

    openai_revised = openai_initial
    claude_revised = claude_initial

    if enable_openai and enable_claude and enable_cross_review:
        st.subheader("2. Cross-review revised answers")

        openai_cross_prompt = build_cross_review_prompt(
            user_question=user_question,
            user_context=user_context,
            attachment_context=attachment_context,
            own_initial_answer=openai_initial,
            other_model_name="Claude",
            other_initial_answer=claude_initial,
        )

        claude_cross_prompt = build_cross_review_prompt(
            user_question=user_question,
            user_context=user_context,
            attachment_context=attachment_context,
            own_initial_answer=claude_initial,
            other_model_name="OpenAI",
            other_initial_answer=openai_initial,
        )

        with st.spinner("Letting the models review each other..."):
            openai_revised = call_openai(openai_cross_prompt, openai_model)
            claude_revised = call_claude(claude_cross_prompt, claude_model)

        col3, col4 = st.columns(2)
        with col3:
            st.markdown("### OpenAI revised after seeing Claude")
            st.markdown(openai_revised)
        with col4:
            st.markdown("### Claude revised after seeing OpenAI")
            st.markdown(claude_revised)

    if enable_final_synthesis:
        st.subheader("3. Final answer")

        if enable_openai and enable_claude:
            final_prompt = build_final_synthesis_prompt(
                user_question=user_question,
                user_context=user_context,
                openai_initial=openai_initial,
                claude_initial=claude_initial,
                openai_revised=openai_revised,
                claude_revised=claude_revised,
            )

            with st.spinner("Creating final synthesis with OpenAI..."):
                final_answer = call_openai(final_prompt, openai_model)
        elif enable_openai:
            final_answer = openai_initial
        else:
            final_answer = claude_initial

        st.markdown(final_answer)

        st.download_button(
            label="Download final answer as TXT",
            data=final_answer,
            file_name="multi_model_final_answer.txt",
            mime="text/plain",
        )

    with st.expander("Debug: extracted attachment text"):
        st.text_area("Attachment context sent to models", attachment_context, height=300)
